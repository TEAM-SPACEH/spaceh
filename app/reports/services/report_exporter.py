"""R-05 보고서 파일 내보내기: Markdown·DOCX 변환, 워터마크, 다운로드(통제 확인·감사 로그)."""

# ─────────────────────────────────────────────────────────────
# services 파일 = "주방"이에요. 화면 코드(st.*)는 쓰지 않아요. (팀 규칙)
#
# 다운로드 흐름 (명세서 export_report) — 버튼을 "누른 순간"에만 실행돼요
#   ① 다운로드 통제 확인 (A-06 decide_report_download) → 차단이면 PermissionError
#      (차단 감사 로그는 A-06이 남겨요. 여기서 또 남기면 2건이 돼요)
#   ② 선택한 형식으로 파일 만들기 → to_markdown() / to_docx()
#   ③ A-06이 준 워터마크 문구를 apply_watermark()로 넣어요
#   ④ 허용된 다운로드를 감사 로그(A-04)에 남겨요
#
# 파일 안 순서 (Markdown·DOCX 똑같아요)
#   제목·조건 → 1페이지 요약 → 상세보고서 → 각주
#   요약 항목마다 [상세 보기] 링크 → 같은 파일 안 상세 섹션으로 이동해요.
# ─────────────────────────────────────────────────────────────

import html
import io
import logging
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.admin.services.audit_service import record_audit_log
from app.reports.models.report_models import Report, ReportTemplate
from app.reports.repositories.report_repository import find_by_id
from app.reports.services.citation_mapper import build_section_anchor
from app.reports.services.download_policy_service import decide_report_download
from app.reports.services.prompts import NO_EVIDENCE_TEXT, SUMMARY_KIND_TITLES
from app.reports.services.report_generator import get_company_names

logger = logging.getLogger(__name__)

# 파일 형식 (명세서: file_format은 'md' | 'docx')
FORMAT_MARKDOWN = "md"
FORMAT_DOCX = "docx"
# 브라우저에 "이 파일은 어떤 종류인지" 알려 주는 값(MIME 타입)이에요. 다운로드 버튼에 넘겨요.
MIME_TYPES = {
    FORMAT_MARKDOWN: "text/markdown",
    FORMAT_DOCX: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}

# 양식(영어 값) → 파일에 보일 한글 이름
TEMPLATE_TITLES = {
    ReportTemplate.COMPREHENSIVE: "종합 보고서",
    ReportTemplate.QUARTERLY: "분기 실적 보고서",
}

# 파일에 들어가는 고정 문구
SUMMARY_HEADING = "1페이지 요약"
DETAIL_HEADING = "상세보고서"
FOOTNOTE_HEADING = "각주"
NOTES_HEADING = "유의사항"
DETAIL_LINK_TEXT = "상세 보기"
NO_SUMMARY_TEXT = "1페이지 요약이 없습니다."
NO_FOOTNOTE_TEXT = "표시할 각주가 없습니다."
PARTIAL_TEXT = "일부 문장이 근거 검증에 실패해 제외되었습니다."

# 감사 로그(A-04) 행위 코드
AUDIT_DOWNLOAD = "download"

# 파일 이름에 쓸 수 없는 글자 (Windows 기준: \ / : * ? " < > |)
FILE_NAME_BAD_CHARS = re.compile(r'[\\/:*?"<>|\s]+')


# ─────────────────────────────────────────────────────────────
# ① 다운로드 통제 (A-06)
# ─────────────────────────────────────────────────────────────
def get_download_policy(user: dict, report: Report, file_format: str) -> dict:
    """A-06 판정을 부른다. 근거 문서를 하나라도 못 보는 사용자면 차단된다."""
    # 판정 자체는 이승준님 decide_report_download()가 해요. 여기서는 보고서에서 필요한 값만 꺼내 넘겨요.
    # 돌려받는 모양: {"allowed": True/False, "reason": 차단 사유, "watermark": 넣을 문구, "blocked_doc_ids": [...]}
    # ⚠️ 차단이면 A-06이 감사 로그를 남겨요. 그래서 화면을 그릴 때(버튼 표시)는 부르지 않고,
    #    버튼을 누른 순간에만 불러요. 화면은 다시 실행될 때마다 그려져서, 그때 부르면 차단 기록이 계속 쌓여요.
    return decide_report_download(
        user, _doc_ids(report), file_format, _report_id_text(report),
        max_access_level=report.max_access_level,  # 기밀이 있으면 A-06이 차단해요 (docs/decisions.md 10/30)
    )


# ─────────────────────────────────────────────────────────────
# ② 다운로드 (명세서 export_report)
# ─────────────────────────────────────────────────────────────
def export_report(user: dict, report_id: int, file_format: str) -> bytes:
    """R-05 저장된 본인 보고서를 찾아, 통제를 다시 확인하고 선택한 형식의 파일을 만든다."""
    # find_by_id는 보고서 ID와 작성자 ID를 함께 걸어서 남의 보고서는 None이에요.
    report = find_by_id(report_id, user.get("id"))
    if report is None:
        raise LookupError("보고서를 찾을 수 없습니다.")
    return export_report_file(user, report, file_format)


def export_report_file(user: dict, report: Report, file_format: str) -> bytes:
    """R-05 통제 확인 → 파일 만들기 → 워터마크 → 감사 로그. 차단이면 PermissionError를 낸다."""
    if file_format not in MIME_TYPES:  # 'md', 'docx' 말고는 받지 않아요
        raise ValueError(f"지원하지 않는 형식입니다: {file_format}")

    # ① 내려받기 직전에 통제를 확인해요. 화면을 거치지 않은 요청도 여기서 막혀요.
    policy = get_download_policy(user, report, file_format)
    if not policy["allowed"]:
        # 빈 파일을 돌려주지 않고 오류를 내요. 다운로드 버튼이 실패로 끝나서
        # 사용자가 "받아졌다"고 오해하지 않아요. (#109 원문 다운로드와 같은 방식)
        # 차단 감사 로그는 A-06이 이미 남겼어요.
        raise PermissionError(policy["reason"])

    # ② 형식에 맞는 파일 만들기. Markdown은 글자라서 bytes로 바꿔요(utf-8 = 한글이 깨지지 않는 방식).
    if file_format == FORMAT_MARKDOWN:
        content = to_markdown(report).encode("utf-8")
    else:
        content = to_docx(report)

    # ③ 워터마크 문구는 A-06 한 곳에서 만들어요. 여기서는 받은 문구를 그대로 넣기만 해요.
    if policy["watermark"]:
        content = apply_watermark(content, file_format, policy["watermark"])

    # ④ 파일을 다 만든 뒤에 기록해요. 만들다 실패하면 "내려받음"으로 남지 않게요.
    _record_audit(user, report, file_format)
    return content


def build_file_name(report: Report, file_format: str) -> str:
    """다운로드 파일 이름을 만든다. 예: 스마트_물류_종합_보고서.md"""
    name = f"{report.request.topic} {TEMPLATE_TITLES[report.request.template]}"
    # 파일 이름에 쓸 수 없는 글자와 띄어쓰기를 "_"로 바꿔요. 앞뒤 "_"는 떼요.
    return f"{FILE_NAME_BAD_CHARS.sub('_', name).strip('_')}.{file_format}"


def _doc_ids(report: Report) -> list[str]:
    """각주의 근거 문서 ID 목록 (순서 유지, 중복 제거)."""
    # dict.fromkeys: 순서를 지키면서 같은 값을 한 번만 남기는 방법이에요.
    return list(dict.fromkeys(note.doc_id for note in report.footnotes))


def _report_id_text(report: Report) -> str | None:
    """감사 로그·A-06에 넘길 보고서 ID 글자. 저장 전 보고서면 None."""
    return str(report.report_id) if report.report_id is not None else None


def _record_audit(user: dict, report: Report, file_format: str) -> None:
    """허용된 보고서 다운로드를 감사 로그(A-04)에 1건 남긴다. 기록이 실패해도 다운로드는 계속한다."""
    # 차단 기록은 A-06(decide_report_download)이 남기므로 여기서는 허용만 남겨요.
    companies = list(dict.fromkeys(note.company_id for note in report.footnotes))
    try:
        record_audit_log(
            user.get("id"),          # 행동한 사용자 (users.id)
            AUDIT_DOWNLOAD,          # 행위 코드: download
            _report_id_text(report), # 대상 보고서 ID (DB 칸 길이 100자 안이라 안전해요)
            None,                    # ip_address: 저장하지 않는 칸이라 None (A-04 규칙)
            False,                   # 허용된 다운로드
            None,                    # 차단 사유 없음
            target_type="report",
            doc_ids=_doc_ids(report),
            doc_companies=companies,
            context_detail={"file_format": file_format},  # 어떤 형식으로 받았는지
        )
    except Exception:
        # DB 오류 등으로 기록이 실패해도 다운로드는 막지 않아요. (A-04 공통 규칙)
        logger.exception("감사 로그 기록 실패(보고서 다운로드) report_id=%s", report.report_id)


# ─────────────────────────────────────────────────────────────
# ③ Markdown 변환
# ─────────────────────────────────────────────────────────────
def to_markdown(report: Report) -> str:
    """R-05 제목·조건 → 1페이지 요약 → 상세보고서 → 각주 순서로 하나의 Markdown 문서를 만든다."""
    lines: list[str] = []  # 한 줄씩 모았다가 마지막에 줄바꿈으로 이어 붙여요.

    # ── 제목·조건 ──
    lines += [f"# {_md(_report_title(report))}", ""]
    lines += [f"- {_md(label)}: {_md(value)}" for label, value in _condition_rows(report)]
    lines.append("")

    # ── 1페이지 요약 ──
    lines += [f"## {SUMMARY_HEADING}", ""]
    summary = report.exec_summary
    if summary is None:
        lines += [NO_SUMMARY_TEXT, ""]
    else:
        names = _company_name_map(report)
        for kind, title in SUMMARY_KIND_TITLES.items():  # 핵심 내용 → 계열사별 시사점 → 권고안
            lines += [f"### {title}", ""]
            for item in (item for item in summary.items if item.kind == kind):
                # 계열사별 시사점이면 앞에 회사 이름을 굵게 붙여요.
                company = f"**{_md(names.get(item.company_id, item.company_id))}** " if item.company_id else ""
                # [상세 보기](#sec-status): 같은 파일 안 <a id="sec-status">로 이동하는 링크예요.
                link = f"[{DETAIL_LINK_TEXT}](#{build_section_anchor(item.section_ref)})"
                lines.append(f"- {company}{_md(item.text)} {link}")
            lines.append("")
        if summary.notes:  # 코드가 쓴 유의사항
            lines += [f"### {NOTES_HEADING}", ""]
            lines += [f"- {_md(note)}" for note in summary.notes]
            lines.append("")

    # 가로줄(---)로 요약과 상세보고서를 나눠요. (DOCX에서는 페이지 나누기)
    lines += ["---", "", f"## {DETAIL_HEADING}", ""]

    # ── 상세보고서 ──
    for section in report.sections:
        # 섹션 제목이 한글이라 자동 링크 주소가 뷰어마다 달라요.
        # 그래서 <a id="sec-status"></a>로 주소를 직접 달아 요약의 [상세 보기]가 항상 찾아오게 해요.
        lines += [f'<a id="{build_section_anchor(section.section_type)}"></a>', ""]
        lines += [f"### {_md(section.title)}", ""]
        if not section.has_evidence:
            lines += [NO_EVIDENCE_TEXT, ""]
            continue
        if section.partial:
            lines += [f"> {PARTIAL_TEXT}", ""]  # > 는 Markdown 인용(회색 박스) 표시예요
        lines += [_md(section.body), ""]

    # ── 각주 ──
    lines += [f"## {FOOTNOTE_HEADING}", ""]
    if not report.footnotes:
        lines += [NO_FOOTNOTE_TEXT, ""]
    for note in report.footnotes:
        # 예: - [1] 현대엘리베이터 · elv_2025_report.pdf · p.12
        lines.append(f"- [{note.number}] {_md(_footnote_source(note))}")
        if note.quote:
            # 줄 끝 공백 2칸 없이 들여쓰기한 줄은 같은 항목 안의 다음 문단으로 보여요.
            lines += ["", f"  “{_md(note.quote)}”"]
    lines.append("")
    return "\n".join(lines)


def _md(text: str) -> str:
    """AI·문서에서 온 글자를 Markdown에 안전하게 넣는다."""
    # < > & 를 &lt; &gt; &amp; 로 바꿔요. 뷰어가 본문 속 <script> 같은 글자를 HTML로 실행하지 않게요.
    # quote=False: 따옴표는 그대로 둬요(본문에 자주 나오고, 바꿔도 위험이 줄지 않아서).
    return html.escape(text, quote=False)


# ─────────────────────────────────────────────────────────────
# ④ DOCX 변환
# ─────────────────────────────────────────────────────────────
def to_docx(report: Report) -> bytes:
    """R-05 Markdown과 같은 순서로 Word 파일을 만든다. 요약 뒤 페이지 나누기, 섹션 책갈피, 내부 링크 포함."""
    document = Document()  # python-docx: 빈 Word 문서
    _set_korean_font(document)

    # ── 제목·조건 ──
    document.add_heading(_report_title(report), level=0)  # level=0 = 문서 제목 스타일
    for label, value in _condition_rows(report):
        document.add_paragraph(f"{label}: {value}")

    # ── 1페이지 요약 ──
    document.add_heading(SUMMARY_HEADING, level=1)
    summary = report.exec_summary
    if summary is None:
        document.add_paragraph(NO_SUMMARY_TEXT)
    else:
        names = _company_name_map(report)
        for kind, title in SUMMARY_KIND_TITLES.items():
            document.add_heading(title, level=2)
            for item in (item for item in summary.items if item.kind == kind):
                # "List Bullet" = Word 기본 글머리 기호(•) 스타일
                paragraph = document.add_paragraph(style="List Bullet")
                if item.company_id:
                    paragraph.add_run(f"{names.get(item.company_id, item.company_id)} ").bold = True
                paragraph.add_run(f"{item.text} ")
                # 요약 항목 → 상세 섹션 책갈피로 가는 문서 안 링크
                _add_internal_link(paragraph, DETAIL_LINK_TEXT, _bookmark_name(item.section_ref))
        if summary.notes:
            document.add_heading(NOTES_HEADING, level=2)
            for note in summary.notes:
                document.add_paragraph(note, style="List Bullet")

    # 명세서: 1페이지 요약 뒤에 페이지 나누기 → 상세보고서는 다음 장부터 시작해요.
    document.add_page_break()

    # ── 상세보고서 ──
    document.add_heading(DETAIL_HEADING, level=1)
    for bookmark_id, section in enumerate(report.sections):
        heading = document.add_heading(section.title, level=2)
        # 섹션 제목에 책갈피를 달아요. 요약의 [상세 보기]가 여기로 이동해요.
        _add_bookmark(heading, _bookmark_name(section.section_type), bookmark_id)
        if not section.has_evidence:
            document.add_paragraph(NO_EVIDENCE_TEXT)
            continue
        if section.partial:
            document.add_paragraph(PARTIAL_TEXT).runs[0].italic = True
        # 빈 줄로 나뉜 문단(현황의 회사별 문단 등)을 Word 문단으로 하나씩 넣어요.
        for block in section.body.split("\n\n"):
            if block.strip():
                document.add_paragraph(block.strip())

    # ── 각주 ──
    document.add_heading(FOOTNOTE_HEADING, level=1)
    if not report.footnotes:
        document.add_paragraph(NO_FOOTNOTE_TEXT)
    for note in report.footnotes:
        paragraph = document.add_paragraph()
        paragraph.add_run(f"[{note.number}] ").bold = True
        paragraph.add_run(_footnote_source(note))
        if note.quote:
            quote = paragraph.add_run(f"\n“{note.quote}”")  # 인용은 다음 줄에 회색 기울임으로
            quote.italic = True
            quote.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # 메모리 안의 가짜 파일(BytesIO)에 저장한 뒤 bytes로 꺼내요. (디스크에 파일을 만들지 않아요)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _set_korean_font(document) -> None:
    """기본 글꼴을 맑은 고딕으로 바꾼다. (Word는 한글 글꼴을 따로 지정해야 해요)"""
    style = document.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(10.5)
    # font.name은 영문 글꼴만 바꿔요. 한글(eastAsia) 글꼴은 XML에 직접 넣어야 적용돼요.
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")


def _bookmark_name(section_ref) -> str:
    """섹션 앵커("sec-status")를 Word 책갈피 이름("sec_status")으로 바꾼다."""
    # Word 책갈피 이름에는 "-"를 쓸 수 없어서 "_"로 바꿔요. 화면·Markdown과 같은 앵커에서 출발해요.
    return build_section_anchor(section_ref).replace("-", "_")


def _add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    """문단에 Word 책갈피를 단다."""
    # python-docx에는 책갈피 기능이 없어서 Word XML(w:bookmarkStart / w:bookmarkEnd)을 직접 넣어요.
    # 시작·끝 표시가 같은 id로 짝을 이뤄요.
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    # 문단 설정(pPr) 바로 뒤에 시작, 글자 뒤에 끝을 넣어 제목 전체를 감싸요.
    p = paragraph._p
    p.insert(1 if p.pPr is not None else 0, start)
    p.append(end)


def _add_internal_link(paragraph, text: str, bookmark: str) -> None:
    """문단 끝에 같은 문서 안 책갈피로 가는 링크(파란 밑줄 글자)를 붙인다."""
    # w:hyperlink에 w:anchor(책갈피 이름)를 주면 웹 주소가 아니라 문서 안으로 이동해요.
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), bookmark)

    run = OxmlElement("w:r")          # 글자 덩어리(run)
    run_style = OxmlElement("w:rPr")  # 글자 모양: 파란색 + 밑줄 (링크처럼 보이게)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_style.append(color)
    run_style.append(underline)
    run.append(run_style)

    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    link.append(run)
    paragraph._p.append(link)


# ─────────────────────────────────────────────────────────────
# ⑤ 워터마크
# ─────────────────────────────────────────────────────────────
def apply_watermark(content: bytes, file_format: str, text: str) -> bytes:
    """R-05 보고서 파일에 다운로드한 사용자와 일시를 담은 워터마크를 넣는다. 본문·각주는 바꾸지 않는다."""
    # 📌 명세서 입력은 user인데, 리뷰 합의로 문구는 A-06 한 곳에서 만들기로 해서 완성된 문구(text)를 받아요.
    #    (A-06 decide_report_download()의 "watermark" 값. 사용자·일시가 들어 있어요)
    if file_format == FORMAT_MARKDOWN:
        # Markdown: 맨 위와 맨 아래에 인용(>) 줄로 넣어요. 본문은 그대로 사이에 둬요.
        body = content.decode("utf-8")
        return f"> {text}\n\n{body}\n> {text}\n".encode("utf-8")
    if file_format == FORMAT_DOCX:
        # DOCX: 모든 페이지 머리글에 회색 글자로 넣어요. 머리글은 본문과 따로라 본문이 바뀌지 않아요.
        document = Document(io.BytesIO(content))
        for section in document.sections:  # Word "구역"마다 머리글이 따로 있어요
            header = section.header.paragraphs[0]
            run = header.add_run(text)
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()
    raise ValueError(f"지원하지 않는 형식입니다: {file_format}")


# ─────────────────────────────────────────────────────────────
# Markdown·DOCX 공통 도우미
# ─────────────────────────────────────────────────────────────
def _report_title(report: Report) -> str:
    """보고서 제목. 예: 스마트 물류 종합 보고서"""
    return f"{report.request.topic} {TEMPLATE_TITLES[report.request.template]}"


def _condition_rows(report: Report) -> list[tuple[str, str]]:
    """제목 아래에 넣을 생성 조건 (이름, 값) 목록."""
    request = report.request
    return [
        ("양식", TEMPLATE_TITLES[request.template]),
        ("대상 회사", ", ".join(get_company_names(request.company_ids))),
        ("기간", f"{request.start_date:%Y-%m-%d} ~ {request.end_date:%Y-%m-%d}"),
    ]


def _company_name_map(report: Report) -> dict[str, str]:
    """회사 코드 → 이름 사전. 계열사별 시사점 앞에 회사 이름을 붙일 때 써요."""
    # zip: 코드 목록과 이름 목록을 같은 순서끼리 짝지어요. 예: {"ELV": "현대엘리베이터"}
    company_ids = report.request.company_ids
    return dict(zip(company_ids, get_company_names(company_ids)))


def _footnote_source(note) -> str:
    """각주 출처 한 줄. 예: 현대엘리베이터 · elv_2025_report.pdf · p.12 (화면 render_footnotes와 같은 형식)"""
    return f"{note.company} · {note.filename} · {note.location}"
