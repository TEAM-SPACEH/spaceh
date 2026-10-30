"""I-01 PPTX 파싱 (슬라이드 번호)."""

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from app.ingestion.parsers.ocr import extract_image_text

from app.ingestion.preprocessing.normalizer import (
    normalize_blocks,
    table_to_markdown,
)


def _iter_shapes(shapes):
    """I-01 그룹 내부까지 재귀적으로 위→아래·왼쪽→오른쪽으로 순회한다."""
    # 저장 순서는 화면의 읽기 순서가 아니라 겹침 순서이므로 좌표로 정렬한다.
    # OOXML에서 생략된 좌표는 0으로 취급해 None과 정수의 비교 오류를 막는다.
    for shape in sorted(shapes, key=lambda item: (item.top or 0, item.left or 0)):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def parse_pptx(file_path: str | Path) -> list[dict]:
    """I-01 슬라이드의 텍스트·표를 위→아래, 왼쪽→오른쪽으로 추출한다."""
    blocks = []
    for number, slide in enumerate(Presentation(file_path).slides, 1):
        title = slide.shapes.title
        section = title.text.strip() if title is not None else ""
        existing = "\n".join(shape.text for shape in _iter_shapes(slide.shapes)
                             if shape.has_text_frame)
        existing += "\n" + "\n".join(
            cell.text for shape in _iter_shapes(slide.shapes) if shape.has_table
            for row in shape.table.rows for cell in row.cells if not cell.is_spanned
        )
        for shape in _iter_shapes(slide.shapes):
            if shape.has_table:
                kind = "table"
                # 병합 시작 셀은 남기고 다른 셀은 비워 원래 열 수를 유지한다.
                rows = [
                    ["" if cell.is_spanned else cell.text for cell in row.cells]
                    for row in shape.table.rows
                ]
                text = table_to_markdown(rows)
            elif shape.has_text_frame:
                is_title = title is not None and shape.shape_id == title.shape_id
                kind = "title" if is_title else "text"
                text = shape.text.replace("\v", "\n")
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                kind = "text"
                text = extract_image_text(shape.image.blob, existing)
                existing += "\n" + text
            else:
                continue
            blocks.append({
                "block_type": kind, "text": text, "page": number,
                "location": (f"슬라이드 {number} 이미지 {shape.shape_id}"
                             if shape.shape_type == MSO_SHAPE_TYPE.PICTURE else f"슬라이드 {number}"),
                "section": section,
            })
    return normalize_blocks(blocks)
