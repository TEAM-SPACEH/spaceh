"""I-01~I-03 생성 문서와 임시 Chroma로 적재 흐름·마스킹·실패 전파를 검증한다."""

import copy
from datetime import datetime
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock, patch

from docx import Document
from pptx import Presentation
from pptx.util import Inches

from app.ingestion.indexing.indexer import DocumentIndexer
from app.ingestion.pipeline import ingest_document
from app.permissions.services.permission_service import apply_rbac_filter


class PipelineTest(unittest.TestCase):
    """임베딩만 대체하고 실제 로더·마스킹·메타데이터·청킹·색인을 연결한다."""

    def setUp(self):
        """업로드 원본과 Chroma를 서비스 데이터와 분리된 임시 경로에 만든다."""
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "sample.docx"
        document = Document()
        document.add_heading("설비 점검", level=1)
        document.add_paragraph("문의 연락처: 010-1234-5678, test@example.com")
        document.add_paragraph("점검 방법 안내 " * 160)
        document.save(self.path)
        self.info = {
            "company_id": "MVX", "company_name": "현대무벡스",
            "file_name": "sample.docx", "file_format": "docx",
            "doc_type": "manual", "title": "", "published_date": None,
            "created_at": datetime(2026, 10, 27, 9),
            "access_level": "internal", "is_group_shared": False,
        }
        self.indexer = DocumentIndexer({"path": Path(self.temp.name) / "chroma"})
        self.addCleanup(self.indexer.close)
        self.embedder = Mock()
        self.embedder.embed_chunks.side_effect = lambda chunks: [[1.0, 0.0] for _ in chunks]

    def run_pipeline(self, **kwargs):
        """공통 입력을 사용하고 필요한 인자만 테스트마다 바꾼다."""
        args = {"file_path": self.path, "doc_id": "pipeline-test",
                "document_info": self.info, "chunking_config": {},
                "embedder": self.embedder, "indexer": self.indexer}
        args.update(kwargs)
        return ingest_document(**args)

    def test_real_flow_masks_before_embedding_and_keeps_input_and_locations(self):
        """임베딩 호출과 저장 내용에 마스킹 전 본문을 넘기지 않고 날짜·권한을 유지한다."""
        before = copy.deepcopy(self.info)
        result = self.run_pipeline()
        self.assertEqual(result["index_status"], "done")
        self.assertEqual(result["error_message"], "")
        chunks = self.embedder.embed_chunks.call_args.args[0]
        self.assertEqual(result["chunk_count"], len(chunks))
        self.assertGreater(len(chunks), 1)
        sent_text = "\n".join(c["text"] for c in chunks)
        self.assertNotIn("010-1234-5678", sent_text)
        self.assertNotIn("test@example.com", sent_text)
        self.assertIn("[마스킹: 연락처]", sent_text)
        self.assertIn("[마스킹: 이메일]", sent_text)
        for c in chunks:
            self.assertEqual(c["metadata"]["published_date"], 20261027)
            self.assertEqual(c["metadata"]["company_id"], "MVX")
            self.assertEqual(c["metadata"]["access_level"], "internal")
            self.assertTrue(c["metadata"]["location"])
        saved = self.indexer.collection.get(include=["documents"])
        self.assertEqual(set(saved["documents"]), {c["text"] for c in chunks})
        self.assertEqual(self.info, before)
        self.assertFalse(self.indexer.client._closed)

    def test_docx_and_pptx_title_pii_is_removed_from_sections_and_table_context(self):
        """제목에 포함된 개인정보가 임베딩 텍스트와 Chroma section에 남지 않는다."""
        title = "담당 010  1234\t5678 test@example.com"
        for extension in ("docx", "pptx"):
            with self.subTest(extension=extension):
                path = Path(self.temp.name) / ("title." + extension)
                if extension == "docx":
                    document = Document()
                    document.add_heading(title, level=1)
                    document.add_paragraph("안전 점검 안내")
                    document.add_table(rows=1, cols=1).cell(0, 0).text = "표 내용"
                    document.save(path)
                else:
                    presentation = Presentation()
                    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
                    slide.shapes.title.text = title
                    slide.shapes.add_table(
                        1, 1, Inches(1), Inches(3), Inches(4), Inches(1),
                    ).table.cell(0, 0).text = "표 내용"
                    presentation.save(path)
                doc_id = "title-" + extension
                info = {**self.info, "file_name": path.name, "file_format": extension}
                result = self.run_pipeline(file_path=path, doc_id=doc_id, document_info=info)
                self.assertEqual(result["index_status"], "done")
                chunks = self.embedder.embed_chunks.call_args.args[0]
                self.assertTrue(any(c["metadata"]["block_type"] == "table" for c in chunks))
                for item in chunks:
                    for value in (item["text"], item["metadata"]["section"]):
                        self.assertNotIn("010-1234-5678", value)
                        self.assertNotIn("010 1234 5678", value)
                        self.assertNotIn("010  1234\t5678", value)
                        self.assertNotIn("test@example.com", value)
                saved = self.indexer.collection.get(
                    where={"doc_id": doc_id}, include=["documents", "metadatas"],
                )
                self.assertEqual(len(saved["ids"]), result["chunk_count"])
                for text, metadata in zip(saved["documents"], saved["metadatas"]):
                    for value in (text, metadata["section"]):
                        self.assertNotIn("010-1234-5678", value)
                        self.assertNotIn("010 1234 5678", value)
                        self.assertNotIn("010  1234\t5678", value)
                        self.assertNotIn("test@example.com", value)
                    self.assertIn("[마스킹: 연락처]", metadata["section"])
                    self.assertIn("[마스킹: 이메일]", metadata["section"])

    def test_registered_display_pii_does_not_reach_embedding_or_storage(self):
        """등록 제목·파일명·파일명 대체 제목을 임베딩과 실제 저장 전에 치환한다."""
        cases = [{"title": "문의 test@example.com"},
                 {"title": "", "file_name": "문의_010  1234  5678.docx"},
                 {"title": "점검 안내", "file_name": "test@example.com.docx"}]
        for index, fields in enumerate(cases):
            with self.subTest(fields=fields):
                info = {**self.info, **fields}
                before = copy.deepcopy(info)
                doc_id = f"display-{index}"
                result = self.run_pipeline(doc_id=doc_id, document_info=info)
                self.assertEqual(result["index_status"], "done")
                chunks = self.embedder.embed_chunks.call_args.args[0]
                saved = self.indexer.collection.get(
                    where={"doc_id": doc_id}, include=["documents", "metadatas"],
                )
                values = [c["text"] for c in chunks] + saved["documents"]
                values.extend(m[key] for m in saved["metadatas"]
                              for key in ("title", "file_name"))
                for value in values:
                    self.assertNotIn("test@example.com", value)
                    self.assertNotIn("010 1234 5678", value)
                    self.assertNotIn("010  1234  5678", value)
                self.assertTrue(all(m["file_name"].endswith(".docx")
                                    for m in saved["metadatas"]))
                self.assertEqual(info, before)

    def test_real_docx_table_keeps_caption_unit_and_source_range(self):
        """실제 DOCX 문단 순번이 바뀌어도 표의 제목·단위와 출처 범위를 유지한다."""
        document = Document()
        document.add_heading("재무 현황", level=1)
        document.add_paragraph("표 1. 매출 내역")
        document.add_paragraph("(단위: 천원)")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "매출 123"
        document.save(self.path)
        self.assertEqual(self.run_pipeline()["index_status"], "done")
        chunks = self.embedder.embed_chunks.call_args.args[0]
        table = next(c for c in chunks if c["metadata"]["block_type"] == "table")
        self.assertIn("표 1. 매출 내역", table["text"])
        self.assertIn("(단위: 천원)", table["text"])
        self.assertEqual(table["metadata"]["location"], "문단 2 ~ 표 4")

    def test_real_docx_intro_is_not_split_at_each_paragraph(self):
        """짧은 DOCX 소개서의 같은 절을 한 청크로 묶고 문단별 위치 범위는 보존한다."""
        document = Document()
        document.add_heading("회사 소개", level=1)
        document.add_paragraph("물류 자동화 안내")
        document.add_paragraph("설비 유지보수 안내")
        document.save(self.path)
        result = self.run_pipeline(document_info={**self.info, "doc_type": "intro"})
        self.assertEqual(result["index_status"], "done")
        self.assertEqual(result["chunk_count"], 1)
        item = self.embedder.embed_chunks.call_args.args[0][0]
        self.assertEqual(item["metadata"]["location"], "문단 1 ~ 문단 3")

    def test_saved_documents_respect_company_permissions(self):
        """저장한 자사 internal 문서를 자사에만 보이고 타사 검색 후보에서는 제외한다."""
        self.assertEqual(self.run_pipeline()["index_status"], "done")
        with patch("app.permissions.services.permission_service.get_approved_doc_ids",
                   return_value=[]):
            for company, expected in (("MVX", True), ("ELV", False)):
                where = apply_rbac_filter(
                    {"id": 1, "role": "employee", "company_id": company}, None)
                found = self.indexer.collection.query(
                    query_embeddings=[[1.0, 0.0]], where=where,
                    n_results=1, include=["metadatas"],
                )
                self.assertEqual(bool(found["ids"][0]), expected)

    def test_failures_stop_later_stages_and_hide_exception_text(self):
        """각 단계의 예외는 실패 결과가 되며 원본·키·경로가 오류 정보로 새지 않는다."""
        steps = [("load_document", "파일 로딩"),
                 ("build_metadata", "메타데이터 구성"),
                 ("mask_sensitive_data", "개인정보 마스킹"),
                 ("split_into_chunks", "청킹")]
        for function, stage in steps:
            with self.subTest(stage=stage):
                self.embedder.reset_mock()
                with patch("app.ingestion.pipeline." + function,
                           side_effect=RuntimeError("private document / fake-secret")):
                    result = self.run_pipeline()
                self.assertEqual(result["index_status"], "failed")
                self.assertEqual(result["chunk_count"], 0)
                self.assertEqual(result["error_message"], stage + " 단계 실패 (RuntimeError)")
                self.embedder.embed_chunks.assert_not_called()
        self.assertEqual(self.indexer.collection.count(), 0)

    def test_embedding_failure_does_not_write_and_can_be_retried(self):
        """임베딩 실패 시 색인하지 않고 주입된 같은 객체로 재시도할 수 있다."""
        self.embedder.embed_chunks.side_effect = RuntimeError("private request body")
        result = self.run_pipeline()
        self.assertEqual(result["error_message"], "임베딩 단계 실패 (RuntimeError)")
        self.assertEqual(self.indexer.collection.count(), 0)
        self.embedder.embed_chunks.side_effect = lambda chunks: [[1.0, 0.0] for _ in chunks]
        self.assertEqual(self.run_pipeline()["index_status"], "done")

    def test_indexer_failure_or_wrong_success_is_not_reported_as_done(self):
        """저장 실패나 다른 문서·개수의 응답을 완료 상태로 반환하지 않는다."""
        failed = Mock()
        failed.index_chunks.side_effect = RuntimeError("private storage details")
        self.assertEqual(self.run_pipeline(indexer=failed)["error_message"],
                         "색인 저장 단계 실패 (RuntimeError)")
        for reply in ({"doc_id": "other", "index_status": "done", "chunk_count": 2},
                      {"doc_id": "pipeline-test", "index_status": "failed", "chunk_count": 0},
                      {"doc_id": "pipeline-test", "index_status": "done", "chunk_count": 999}):
            with self.subTest(reply=reply):
                wrong = Mock()
                wrong.index_chunks.return_value = reply
                self.assertEqual(self.run_pipeline(indexer=wrong)["index_status"], "failed")

    def test_foreign_chunk_collision_returns_failure_and_preserves_owner(self):
        """소유권 충돌을 파이프라인 실패로 반환하고 다른 문서의 원본을 보존한다."""
        foreign_id = "pipeline-test#1#0001"
        metadata = {"doc_id": "pipeline-test#1", "company_id": "MVX",
                    "access_level": "internal", "is_group_shared": False}
        foreign = {"chunk_id": foreign_id, "text": "보존할 다른 문서",
                   "metadata": metadata}
        self.indexer.index_chunks("pipeline-test#1", [foreign], [[1.0, 0.0]])
        conflicting = {**foreign, "text": "덮어쓰기 시도",
                       "metadata": {**metadata, "doc_id": "pipeline-test"}}
        with patch("app.ingestion.pipeline.split_into_chunks", return_value=[conflicting]):
            result = self.run_pipeline()
        self.assertEqual(result["index_status"], "failed")
        self.assertEqual(result["chunk_count"], 0)
        self.assertEqual(result["error_message"], "색인 저장 단계 실패 (ValueError)")
        saved = self.indexer.collection.get(ids=[foreign_id])
        self.assertEqual(saved["documents"], [foreign["text"]])
        self.assertEqual(saved["metadatas"], [metadata])

    def test_invalid_metadata_or_config_does_not_reach_embedding(self):
        """형식 불일치·권한 오류·잘못된 청킹 설정을 API 요청 전에 차단한다."""
        for info in ({**self.info, "file_format": "pdf"},
                     {**self.info, "is_group_shared": True, "access_level": "confidential"}):
            self.assertEqual(self.run_pipeline(document_info=info)["index_status"], "failed")
        self.assertEqual(self.run_pipeline(chunking_config={"chunk_size": 0})["index_status"],
                         "failed")
        self.embedder.embed_chunks.assert_not_called()
        self.assertEqual(self.indexer.collection.count(), 0)


if __name__ == "__main__":
    unittest.main()
