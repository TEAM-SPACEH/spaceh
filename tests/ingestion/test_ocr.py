"""I-01 OCR의 위치·중복·실패 처리. 외부 모델 호출은 대체한다."""

import io
import base64
import json
import sys
from types import SimpleNamespace
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import MagicMock, patch

from PIL import Image
from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn as docx_qn
from pptx import Presentation
from pptx.util import Inches
from pptx.oxml.ns import qn as pptx_qn

from app.ingestion.parsers import ocr
from app.ingestion.parsers.docx import parse_docx, _paragraph_text
from app.ingestion.parsers.pptx import parse_pptx


def image_bytes():
    with Image.new("RGB", (100, 100), "white") as image, io.BytesIO() as buffer:
        image.save(buffer, format="PNG")
        return buffer.getvalue()


class OcrTest(unittest.TestCase):
    @patch("app.ingestion.parsers.docx.extract_image_text")
    def test_docx_textbox_and_fallback_are_excluded(self, extract):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "textbox.docx"
            doc = Document()
            paragraph = doc.add_paragraph("본문 문장")
            paragraph._p.append(parse_xml('''
                <mc:AlternateContent
                  xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
                  xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
                  xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
                  xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
                  <mc:Choice Requires="a"><w:r><w:drawing><w:txbxContent>
                    <w:p><w:r><w:t>상자 글자</w:t><a:blip r:embed="rIdMissing"/></w:r></w:p>
                  </w:txbxContent></w:drawing></w:r></mc:Choice>
                  <mc:Fallback><w:r><w:t>상자 글자</w:t><a:blip r:embed="rIdMissing"/></w:r></mc:Fallback>
                </mc:AlternateContent>'''))
            doc.save(path)
            self.assertEqual([b["text"] for b in parse_docx(path)], ["본문 문장"])
            extract.assert_not_called()

    def test_worker_library_prints_do_not_corrupt_json_response(self):
        class Reader:
            def __init__(self, *args, **kwargs):
                print("library startup")

            def readtext(self, *args, **kwargs):
                print("library diagnostic")
                return ["OCR result"]

        response, diagnostic = io.StringIO(), io.StringIO()
        request = SimpleNamespace(buffer=io.BytesIO(base64.b64encode(b"image") + b"\n"))
        with patch.dict(sys.modules, {"easyocr": SimpleNamespace(Reader=Reader)}), \
                patch.object(sys, "stdin", request), patch.object(sys, "stdout", response), \
                patch.object(sys, "stderr", diagnostic), patch.object(sys, "argv", ["worker", "models"]):
            ocr._worker()
            self.assertIs(sys.stdout, response)
        self.assertEqual(json.loads(response.getvalue()), [True, ["OCR result"]])
        self.assertIn("library startup", diagnostic.getvalue())
        self.assertIn("library diagnostic", diagnostic.getvalue())

    @patch("app.ingestion.parsers.docx.extract_image_text", return_value="정상 그림")
    def test_missing_docx_image_relation_preserves_text_and_other_images(self, extract):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "missing-image.docx"
            doc = Document()
            paragraph = doc.add_paragraph("앞")
            broken = paragraph.add_run()
            broken.add_picture(io.BytesIO(image_bytes()))
            broken._r.xpath(".//a:blip")[0].set(docx_qn("r:embed"), "rIdMissing")
            paragraph.add_run("뒤")
            doc.add_paragraph().add_run().add_picture(io.BytesIO(image_bytes()))
            cell = doc.add_table(1, 1).cell(0, 0)
            cell.text = "셀 본문"
            bad_cell_run = cell.paragraphs[0].add_run()
            bad_cell_run.add_picture(io.BytesIO(image_bytes()))
            bad_cell_run._r.xpath(".//a:blip")[0].set(docx_qn("r:embed"), "rIdMissing")
            doc.save(path)
            with self.assertLogs("app.ingestion.parsers.docx", level="WARNING") as logs:
                blocks = parse_docx(path)
            self.assertEqual([block["page"] for block in blocks], [1, 2, 3])
            self.assertEqual(blocks[0]["text"], "앞뒤")
            self.assertEqual(blocks[1]["text"], "정상 그림")
            self.assertIn("셀 본문", blocks[2]["text"])
            self.assertEqual(len(logs.output), 2)
            extract.assert_called_once()

    def test_docx_tab_stops_are_not_body_tabs(self):
        doc = Document()
        paragraph = doc.add_paragraph("앞")
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(1))
        run = paragraph.add_run("\t뒤\n다음")
        run._r.append(OxmlElement("w:cr"))
        run.add_text("끝")
        self.assertEqual(_paragraph_text(paragraph), "앞\t뒤\n다음\n끝")

    @patch("app.ingestion.parsers.pptx.extract_image_text", return_value="내장 그림")
    def test_pptx_linked_picture_does_not_drop_slide_text(self, extract):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "linked.pptx"
            ppt = Presentation()
            slide = ppt.slides.add_slide(ppt.slide_layouts[6])
            slide.shapes.add_textbox(Inches(0), Inches(0), Inches(2), Inches(1)).text = "본문"
            linked = slide.shapes.add_picture(io.BytesIO(image_bytes()), Inches(1), Inches(1))
            blip = linked._element.xpath(".//a:blip")[0]
            del blip.attrib[pptx_qn("r:embed")]
            blip.set(pptx_qn("r:link"), "external-image")
            slide.shapes.add_picture(io.BytesIO(image_bytes()), Inches(1), Inches(2))
            ppt.save(path)
            with self.assertLogs("app.ingestion.parsers.pptx", level="WARNING"):
                blocks = parse_pptx(path)
            self.assertEqual([block["text"] for block in blocks], ["본문", "내장 그림"])
            extract.assert_called_once()

    def tearDown(self):
        ocr._close_worker()

    @patch("app.ingestion.parsers.pptx.extract_image_text", return_value="정상 그림")
    def test_missing_picture_relation_preserves_other_slides(self, extract):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "missing-relation.pptx"
            ppt = Presentation()
            slide = ppt.slides.add_slide(ppt.slide_layouts[6])
            slide.shapes.add_textbox(Inches(0), Inches(0), Inches(2), Inches(1)).text = "첫 슬라이드"
            broken = slide.shapes.add_picture(io.BytesIO(image_bytes()), Inches(1), Inches(1))
            broken._element.xpath(".//a:blip")[0].set(pptx_qn("r:embed"), "rIdMissing")
            other = ppt.slides.add_slide(ppt.slide_layouts[6])
            other.shapes.add_textbox(Inches(0), Inches(0), Inches(2), Inches(1)).text = "다음 슬라이드"
            other.shapes.add_picture(io.BytesIO(image_bytes()), Inches(1), Inches(1))
            ppt.save(path)
            with self.assertLogs("app.ingestion.parsers.pptx", level="WARNING"):
                blocks = parse_pptx(path)
            self.assertEqual([block["text"] for block in blocks],
                             ["첫 슬라이드", "다음 슬라이드", "정상 그림"])
            self.assertEqual([block["page"] for block in blocks], [1, 2, 2])
            extract.assert_called_once()

    @patch.object(ocr, "_process")
    @patch.object(ocr, "_connection")
    def test_duplicate_text_and_empty_lines_are_removed(self, connection, process):
        process.poll.return_value = None
        connection.recv.return_value = (True, ["Existing text", "새 글자", "새 글자", " "])
        self.assertEqual(ocr.extract_image_text(image_bytes(), "Existing text"), "새 글자")

    @patch.object(ocr, "_process")
    @patch.object(ocr, "_connection")
    def test_timeout_stops_worker(self, connection, process):
        process.poll.return_value = None
        connection.poll.return_value = False
        with self.assertLogs(ocr.logger, level="WARNING"):
            self.assertEqual(ocr.extract_image_text(image_bytes()), "")
        process.terminate.assert_called_once()

    @patch.object(ocr, "_process")
    @patch.object(ocr, "_connection")
    def test_model_failure_is_logged_and_skipped(self, connection, process):
        process.poll.return_value = None
        connection.recv.return_value = (False, [])
        with self.assertLogs(ocr.logger, level="WARNING"):
            self.assertEqual(ocr.extract_image_text(image_bytes()), "")

    @patch.object(ocr, "_process")
    @patch.object(ocr, "_connection")
    def test_short_numeric_labels_are_not_substring_duplicates(self, connection, process):
        process.poll.return_value = None
        connection.recv.return_value = (True, ["10", "%", "매출", "2010"])
        self.assertEqual(ocr.extract_image_text(image_bytes(), "2010\n연간매출액"), "10\n%\n매출")

    @patch.object(ocr, "_process")
    @patch.object(ocr, "_connection")
    def test_broken_pipe_has_safe_model_guidance(self, connection, process):
        process.poll.return_value = None
        connection.send_bytes.side_effect = BrokenPipeError("private")
        with self.assertRaisesRegex(RuntimeError, "모델 준비"):
            ocr._extract_image_text(image_bytes())

    @patch.object(ocr, "_process")
    @patch.object(ocr, "_connection")
    def test_eof_has_safe_model_guidance(self, connection, process):
        process.poll.return_value = None
        connection.recv.side_effect = EOFError("private")
        with self.assertRaisesRegex(ocr.OcrModelError, "모델 준비"):
            ocr._extract_image_text(image_bytes())

    @patch("app.ingestion.parsers.ocr._extract_image_text")
    def test_small_image_is_not_sent_to_worker(self, extract):
        with Image.new("RGB", (99, 200)) as image, io.BytesIO() as buffer:
            image.save(buffer, format="PNG")
            self.assertEqual(ocr.extract_image_text(buffer.getvalue()), "")
        extract.assert_not_called()

    @patch("app.ingestion.parsers.ocr._extract_image_text")
    def test_document_time_budget_skips_worker(self, extract):
        from core import config
        @ocr.ocr_document
        def parse():
            return ocr.extract_image_text(image_bytes())
        with patch.object(config, "OCR_DOCUMENT_SECONDS", 0), self.assertLogs(ocr.logger):
            self.assertEqual(parse(), "")
        extract.assert_not_called()

    @patch("app.ingestion.parsers.ocr._extract_image_text", side_effect=ocr.OcrModelError("model"))
    def test_missing_model_preserves_native_docx_text(self, extract):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "native.docx"
            doc = Document()
            doc.add_paragraph("기존 텍스트")
            doc.add_picture(io.BytesIO(image_bytes()))
            doc.save(path)
            with self.assertLogs(ocr.logger):
                self.assertEqual([b["text"] for b in parse_docx(path)], ["기존 텍스트"])

    @patch("app.ingestion.parsers.ocr._extract_image_text", return_value="OCR")
    def test_document_count_budget_is_enforced(self, extract):
        from core import config
        @ocr.ocr_document
        def parse():
            return [ocr.extract_image_text(image_bytes()) for _ in range(3)]
        with patch.object(config, "OCR_MAX_IMAGES", 1), self.assertLogs(ocr.logger, level="WARNING"):
            self.assertEqual(parse(), ["OCR", "", ""])
        extract.assert_called_once()

    @patch("app.ingestion.parsers.docx.extract_image_text", return_value="이미지 글자")
    def test_docx_paragraph_and_table_images_keep_location(self, extract):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "sample.docx"
            doc = Document()
            doc.add_paragraph("본문").add_run().add_picture(io.BytesIO(image_bytes()))
            doc.add_table(1, 1).cell(0, 0).paragraphs[0].add_run().add_picture(io.BytesIO(image_bytes()))
            doc.save(path)
            blocks = parse_docx(path)
            self.assertEqual(blocks[0]["location"], "문단 1")
            self.assertEqual(blocks[1]["location"], "표 2")
            self.assertTrue(all("이미지 글자" in b["text"] for b in blocks))
            self.assertEqual(extract.call_count, 2)

    @patch("app.ingestion.parsers.docx.extract_image_text", return_value="이미지")
    def test_docx_image_stays_between_text_runs(self, extract):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "order.docx"
            doc = Document()
            paragraph = doc.add_paragraph("앞")
            paragraph.add_run().add_picture(io.BytesIO(image_bytes()))
            paragraph.add_run("뒤")
            doc.save(path)
            self.assertEqual(parse_docx(path)[0]["text"], "앞\n이미지\n뒤")

    @patch("app.ingestion.parsers.pptx.extract_image_text", return_value="이미지 글자")
    def test_pptx_image_keeps_slide_and_shape(self, extract):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "sample.pptx"
            ppt = Presentation()
            slide = ppt.slides.add_slide(ppt.slide_layouts[6])
            shape = slide.shapes.add_picture(io.BytesIO(image_bytes()), Inches(1), Inches(1))
            ppt.save(path)
            blocks = parse_pptx(path)
            self.assertEqual(blocks[0]["text"], "이미지 글자")
            self.assertEqual(blocks[0]["location"], f"슬라이드 1 이미지 {shape.shape_id}")
