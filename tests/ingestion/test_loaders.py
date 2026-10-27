"""I-01 생성한 목문서로 순서·위치·표·입력 오류를 검증한다."""

import copy
from pathlib import Path
import tempfile
import unittest
from unittest.mock import MagicMock, patch

from docx import Document
from pptx import Presentation
from pptx.util import Inches

from app.ingestion.parsers.loader import load_document
from app.ingestion.parsers.pdf import parse_pdf
from app.ingestion.preprocessing.normalizer import (
    normalize_blocks,
    table_to_markdown,
)


def write_pdf(path: Path, streams: list[bytes]) -> None:
    """외부 PDF 생성 패키지 없이 텍스트·도형 목문서를 만든다."""
    pages = " ".join(f"{4 + i * 2} 0 R" for i in range(len(streams)))
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Count {len(streams)} /Kids [{pages}] >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for i, stream in enumerate(streams):
        objects.append(
            ("<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
             "/Resources << /Font << /F1 3 0 R >> >> "
             f"/Contents {5 + i * 2} 0 R >>").encode()
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode()
                       + stream + b"\nendstream")
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    start = len(data)
    data.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        data.extend(f"{offset:010} 00000 n \n".encode())
    data.extend((f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
                 f"startxref\n{start}\n%%EOF\n").encode())
    path.write_bytes(data)


class LoaderTest(unittest.TestCase):
    """I-01 로더의 공통 계약과 형식별 추출 결과를 검증한다."""

    def setUp(self):
        """각 테스트의 목문서를 격리하고 종료 시 자동 삭제한다."""
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def assert_common_blocks(self, blocks):
        """추출된 모든 블록의 필드·내용·원문 위치를 확인한다."""
        for block in blocks:
            self.assertEqual(set(block), {"block_type", "text", "page",
                                          "location", "section"})
            self.assertGreater(block["page"], 0)
            self.assertTrue(block["text"].strip())
            self.assertTrue(block["location"])

    def test_docx_preserves_order_heading_and_merged_cells(self):
        """제목·본문·가로 병합 표 순서와 빈 문단 이후 순번을 검증한다."""
        path = self.root / "sample.DOCX"
        doc = Document()
        doc.add_heading("점검 절차", level=1)
        doc.add_paragraph("  첫 번째   절차  ")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "병합 제목"
        table.cell(1, 0).text = "항목"
        table.cell(1, 1).text = "값"
        doc.add_paragraph("")
        doc.add_paragraph("마지막 절차")
        doc.save(path)
        blocks = load_document(path)
        self.assertEqual([b["block_type"] for b in blocks],
                         ["title", "text", "table", "text"])
        self.assertEqual([b["page"] for b in blocks], [1, 2, 3, 5])
        self.assertEqual(blocks[1]["text"], "첫 번째 절차")
        self.assertEqual(blocks[2]["text"].count("병합 제목"), 1)
        self.assertIn("| 항목 | 값 |", blocks[2]["text"])
        self.assertTrue(all(b["section"] == "점검 절차" for b in blocks))
        self.assert_common_blocks(blocks)

    def test_docx_keeps_nested_table_content(self):
        """부모 셀 안 중첩 표의 내용과 앞뒤 문단 순서를 보존한다."""
        path = self.root / "nested.docx"
        doc = Document()
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "앞"
        cell.add_table(rows=1, cols=1).cell(0, 0).text = "안쪽 표"
        cell.add_paragraph("뒤")
        doc.save(path)
        text = load_document(path)[0]["text"]
        self.assertLess(text.index("앞"), text.index("안쪽 표"))
        self.assertLess(text.index("안쪽 표"), text.index("뒤"))

    def test_docx_vertical_merge_keeps_column_positions(self):
        """세로 병합 연속 셀을 비워 나머지 열이 밀리지 않게 한다."""
        path = self.root / "vertical.docx"
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(1, 0)).text = "세로 병합"
        table.cell(0, 1).text = "위"
        table.cell(1, 1).text = "아래"
        doc.save(path)
        text = load_document(path)[0]["text"]
        self.assertEqual(text.count("세로 병합"), 1)
        self.assertIn("| 세로 병합 | 위 |", text)
        self.assertIn("| | 아래 |", text)

    def test_pptx_reads_visual_order_tables_and_groups(self):
        """슬라이드·그룹의 시각적 순서와 표·제목·위치를 확인한다."""
        path = self.root / "sample.pptx"
        ppt = Presentation()
        slide = ppt.slides.add_slide(ppt.slide_layouts[5])
        slide.shapes.title.text = "제목"
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(4),
                                       Inches(4), Inches(1)).table
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(0, 0).text = "공통"
        table.cell(1, 0).text = "A"
        table.cell(1, 1).text = "B"
        slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4),
                                 Inches(1)).text = "표보다 먼저"
        group_slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        group = group_slide.shapes.add_group_shape()
        group.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = "그룹 내용"
        ppt.save(path)
        blocks = load_document(path)
        self.assertEqual([b["block_type"] for b in blocks],
                         ["title", "text", "table", "text"])
        self.assertEqual(blocks[1]["text"], "표보다 먼저")
        self.assertEqual(blocks[2]["text"].count("공통"), 1)
        self.assertEqual(blocks[3]["page"], 2)
        self.assertEqual(blocks[3]["section"], "")
        self.assert_common_blocks(blocks)

    def test_pptx_table_keeps_soft_breaks_and_paragraphs(self):
        """실제 PPTX 셀의 소프트 줄바꿈·문단을 Markdown에서 보존한다."""
        path = self.root / "line_breaks.pptx"
        ppt = Presentation()
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        table = slide.shapes.add_table(
            1, 1, Inches(1), Inches(1), Inches(4), Inches(2),
        ).table
        table.cell(0, 0).text = "첫 줄\v둘째 줄\n새 문단"
        ppt.save(path)
        blocks = load_document(path)
        self.assertEqual(len(blocks), 1)
        self.assertEqual(blocks[0]["block_type"], "table")
        self.assertIn("첫 줄<br>둘째 줄<br>새 문단", blocks[0]["text"])
        self.assertNotIn("\v", blocks[0]["text"])

    def test_pdf_tables_are_not_duplicated_and_pages_are_preserved(self):
        """PDF 표 문자의 중복을 막고 본문 순서·페이지 위치를 유지한다."""
        path = self.root / "sample.pdf"
        write_pdf(path, [b"""
BT /F1 12 Tf 50 760 Td (Before) Tj ET
50 680 200 40 re S
50 700 m 250 700 l S
150 680 m 150 720 l S
BT /F1 10 Tf 60 705 Td (Item) Tj ET
BT /F1 10 Tf 160 705 Td (Value) Tj ET
BT /F1 10 Tf 60 685 Td (Motor) Tj ET
BT /F1 10 Tf 160 685 Td (42) Tj ET
BT /F1 12 Tf 50 640 Td (After) Tj ET
""", b"BT /F1 12 Tf 50 760 Td (Second page) Tj ET"])
        blocks = load_document(path)
        self.assertEqual([b["block_type"] for b in blocks],
                         ["text", "table", "text", "text"])
        self.assertEqual(blocks[0]["text"], "Before")
        self.assertIn("| Motor | 42 |", blocks[1]["text"])
        self.assertEqual(sum(b["text"].count("Motor") for b in blocks), 1)
        self.assertEqual(blocks[-1]["page"], 2)
        self.assertTrue(blocks[-1]["location"].startswith("p.2"))
        self.assert_common_blocks(blocks)

    def test_scanned_pdf_is_reported_instead_of_silently_skipped(self):
        """이미지뿐인 페이지를 조용히 누락시키지 않고 오류로 알린다."""
        page = MagicMock()
        page.images = [{}]
        page.chars = []
        with patch("app.ingestion.parsers.pdf.pdfplumber.open") as opened:
            opened.return_value.__enter__.return_value.pages = [page]
            with self.assertRaisesRegex(ValueError, "OCR"):
                parse_pdf("scanned.pdf")

    def test_missing_unsupported_corrupt_and_empty_documents(self):
        """잘못된 파일을 성공한 빈 결과로 반환하지 않는지 확인한다."""
        with self.assertRaises(FileNotFoundError):
            load_document(self.root / "missing.pdf")
        path = self.root / "unsupported.xlsx"
        path.write_bytes(b"not an excel file")
        with self.assertRaises(ValueError):
            load_document(path)
        for extension in ("pdf", "docx", "pptx"):
            with self.subTest(extension=extension):
                corrupt = self.root / f"corrupt.{extension}"
                corrupt.write_bytes(b"not a document")
                with self.assertRaises(Exception):
                    load_document(corrupt)
        empty = self.root / "empty.docx"
        Document().save(empty)
        with self.assertRaisesRegex(ValueError, "추출 가능한"):
            load_document(empty)

    def test_oversized_file_is_rejected_before_parsing(self):
        """화면 없는 직접 로더 호출도 상한 초과 파일을 파싱하지 않는다."""
        path = self.root / "oversized.pdf"
        # 큰 파일을 실제로 만들지 않고 제한을 줄여 경계 조건을 검증한다.
        path.write_bytes(b"1234")
        with patch("app.ingestion.parsers.loader.MAX_DOCUMENT_BYTES", 3):
            with patch("app.ingestion.parsers.loader.parse_pdf") as parser:
                with self.assertRaisesRegex(ValueError, "20 MiB"):
                    load_document(path)
                parser.assert_not_called()

    def test_file_at_size_limit_is_passed_to_parser(self):
        """정확히 상한 크기인 파일은 허용하고 기존 블록 반환을 유지한다."""
        path = self.root / "boundary.pdf"
        path.write_bytes(b"123")
        expected = [{"block_type": "text", "text": "허용", "page": 1,
                     "location": "p.1", "section": ""}]
        with patch("app.ingestion.parsers.loader.MAX_DOCUMENT_BYTES", 3):
            with patch("app.ingestion.parsers.loader.parse_pdf",
                       return_value=expected) as parser:
                self.assertEqual(load_document(path), expected)
                parser.assert_called_once_with(path)

    def test_markdown_preserves_cells_and_escapes_special_characters(self):
        """셀의 열 위치·줄바꿈·특수 문자가 Markdown 구조를 깨지 않는다."""
        result = table_to_markdown([
            ["A|B", "line1\nline2"], [None, "x\\y"], ["last"],
        ])
        self.assertIn(r"A\|B", result)
        self.assertIn("line1<br>line2", result)
        self.assertIn(r"x\\y", result)
        self.assertIn("| last |  |", result)
        self.assertEqual(table_to_markdown([[None, ""]]), "")

    def test_normalization_keeps_lines_locations_and_input(self):
        """정규화가 원본·위치·줄바꿈을 유지하며 잘못된 블록을 거부한다."""
        blocks = [{"block_type": "text", "text": " A  B\r\n C\tD ",
                   "page": 1, "location": "문단 1", "section": ""}]
        original = copy.deepcopy(blocks)
        result = normalize_blocks(blocks)
        self.assertEqual(result[0]["text"], "A B\nC D")
        self.assertEqual(blocks, original)
        self.assertEqual(result[0]["location"], "문단 1")
        for key, invalid in [("page", 0), ("page", True), ("text", None),
                             ("block_type", "image"), ("location", "")]:
            with self.subTest(key=key, invalid=invalid):
                with self.assertRaises(ValueError):
                    normalize_blocks([{**blocks[0], key: invalid}])


if __name__ == "__main__":
    unittest.main()
