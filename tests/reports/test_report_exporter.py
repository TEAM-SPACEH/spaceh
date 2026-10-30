"""R-05 보고서 내보내기(report_exporter) 테스트. 저장소 루트에서 `python -m unittest`."""

# ─────────────────────────────────────────────────────────────
# 회사 목록(get_company_options)·감사 로그(record_audit_log)·A-06 판정(decide_report_download)·
# 보고서 조회(find_by_id)는 DB를 써서 가짜로 바꿔 끼워요.
# DOCX는 만든 bytes를 python-docx로 다시 열어서 안을 확인해요. (디스크에 파일을 만들지 않아요)
# ─────────────────────────────────────────────────────────────

import io
import unittest
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn

from app.reports.mock_data import MOCK_COLLECTION, MOCK_REQUEST, make_quote
from app.reports.models.report_models import (
    ExecSummary, Footnote, Report, Section, SectionType, SummaryItem,
)
from app.reports.services import report_exporter
from app.reports.services.prompts import NO_EVIDENCE_TEXT

COMPANY_OPTIONS = "app.reports.services.report_generator.get_company_options"
RECORD_AUDIT_LOG = "app.reports.services.report_exporter.record_audit_log"
DECIDE = "app.reports.services.report_exporter.decide_report_download"
FIND_BY_ID = "app.reports.services.report_exporter.find_by_id"

USER = {"id": 7, "name": "박예원", "employee_no": "E2025001", "role": "employee", "company_id": "ELV"}
WATERMARK = "워터마크: 박예원(E2025001) · 2026-10-29 15:30 다운로드 · 외부 유출 금지"
# A-06 decide_report_download()가 돌려주는 모양 (허용 / 차단)
ALLOWED = {"allowed": True, "reason": None, "watermark": WATERMARK, "blocked_doc_ids": []}
BLOCKED = {"allowed": False, "reason": "열람 권한이 없는 근거 문서가 1건 포함되어 있습니다.",
           "watermark": None, "blocked_doc_ids": ["doc-hri"]}


def setUpModule():
    global _company_patcher
    _company_patcher = patch(COMPANY_OPTIONS, return_value=[
        {"id": "ELV", "name": "현대엘리베이터"},
        {"id": "MVX", "name": "현대무벡스"},
        {"id": "HRI", "name": "현대경제연구원"},
    ])
    _company_patcher.start()


def tearDownModule():
    _company_patcher.stop()


def make_report(exec_summary: ExecSummary | None = None, report_id: int | None = 12) -> Report:
    """보고서 전체 번호가 매겨진 가짜 보고서. 각주 [1]=ELV, [2]=MVX, [3]=HRI. 시너지는 근거 없음."""
    footnotes = [
        Footnote(
            number=number, company=chunk.metadata["company_name"], company_id=chunk.metadata["company_id"],
            filename=chunk.metadata["file_name"], location=chunk.metadata["location"],
            quote=make_quote(chunk.text), text=chunk.text, doc_id=chunk.metadata["doc_id"],
            chunk_id=chunk.chunk_id,
        )
        for number, chunk in enumerate(MOCK_COLLECTION.chunks, start=1)
    ]
    sections = [
        Section(SectionType.OVERVIEW, "개요", "물류 자동화 수요가 늘고 있음.[1][3]"),
        Section(SectionType.STATUS, "현황",
                "현대엘리베이터는 승강기 공급을 늘림.[1]\n\n현대무벡스는 설비 공급을 늘림.[2]", partial=True),
        Section(SectionType.SYNERGY, "시너지", NO_EVIDENCE_TEXT, has_evidence=False),
        Section(SectionType.ADVICE, "제언", "묶음 제안을 검토함.[1][2] <script>alert(1)</script>"),
    ]
    if exec_summary is None:
        exec_summary = ExecSummary(
            items=[
                SummaryItem("key", "물류 자동화 수요가 늘고 있음.[1][3]", SectionType.OVERVIEW),
                SummaryItem("company", "설비 공급을 늘림.[2]", SectionType.STATUS, company_id="MVX"),
                SummaryItem("advice", "묶음 제안을 검토함.[1][2]", SectionType.ADVICE),
            ],
            notes=["시너지 섹션은 근거 자료가 없어 요약에서 제외됨"],
        )
    return Report(request=MOCK_REQUEST, sections=sections, footnotes=footnotes,
                  exec_summary=exec_summary, report_id=report_id)


def docx_texts(content: bytes) -> list[str]:
    """DOCX 본문 문단 글자 목록."""
    return [paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs]


class ToMarkdownTest(unittest.TestCase):

    def setUp(self):
        self.markdown = report_exporter.to_markdown(make_report())

    def test_order(self):
        # 제목·조건 → 1페이지 요약 → 상세보고서 → 각주 순서
        positions = [self.markdown.index(text) for text in (
            "# 스마트 물류 종합 보고서", "## 1페이지 요약", "## 상세보고서", "## 각주",
        )]
        self.assertEqual(positions, sorted(positions))
        self.assertIn("- 대상 회사: 현대엘리베이터, 현대무벡스, 현대경제연구원", self.markdown)
        self.assertIn("- 기간: 2025-01-01 ~ 2025-12-31", self.markdown)

    def test_detail_links_match_anchors(self):
        # 요약의 [상세 보기](#sec-status)가 같은 파일 안 <a id="sec-status">를 가리켜요.
        self.assertIn("[상세 보기](#sec-status)", self.markdown)
        self.assertIn('<a id="sec-status"></a>', self.markdown)
        self.assertIn("- **현대무벡스** 설비 공급을 늘림.[2] [상세 보기](#sec-status)", self.markdown)

    def test_sections_and_footnotes(self):
        self.assertIn(f"### 시너지\n\n{NO_EVIDENCE_TEXT}", self.markdown)
        self.assertIn("> 일부 문장이 근거 검증에 실패해 제외되었습니다.", self.markdown)  # 현황 partial
        self.assertIn("- 시너지 섹션은 근거 자료가 없어 요약에서 제외됨", self.markdown)  # 유의사항
        self.assertIn("- [1] 현대엘리베이터 · elv_2025_report.pdf · p.12", self.markdown)
        self.assertIn("  “물류센터용 고속 화물 승강기", self.markdown)

    def test_html_escaped(self):
        # 본문 속 <script>가 Markdown 뷰어에서 HTML로 실행되지 않게 바꿔요.
        self.assertNotIn("<script>", self.markdown)
        self.assertIn("&lt;script&gt;", self.markdown)

    def test_no_summary(self):
        report = make_report()
        report.exec_summary = None
        self.assertIn("1페이지 요약이 없습니다.", report_exporter.to_markdown(report))


class ToDocxTest(unittest.TestCase):

    def setUp(self):
        self.content = report_exporter.to_docx(make_report())
        self.document = Document(io.BytesIO(self.content))
        self.body = self.document.element.body  # Word XML 본문

    def test_order_and_text(self):
        texts = docx_texts(self.content)
        positions = [texts.index(text) for text in ("스마트 물류 종합 보고서", "1페이지 요약", "상세보고서", "각주")]
        self.assertEqual(positions, sorted(positions))
        self.assertIn("현대엘리베이터는 승강기 공급을 늘림.[1]", texts)  # 빈 줄로 나뉜 문단이 따로 들어가요
        self.assertIn("현대무벡스는 설비 공급을 늘림.[2]", texts)
        self.assertIn(NO_EVIDENCE_TEXT, texts)

    def test_page_break_after_summary(self):
        # 페이지 나누기(w:br type=page)가 "상세보고서" 제목보다 앞에 있어야 해요.
        texts = docx_texts(self.content)
        paragraphs = self.document.paragraphs
        break_index = next(
            index for index, paragraph in enumerate(paragraphs)
            if paragraph._p.findall(f".//{qn('w:br')}[@{qn('w:type')}='page']")
        )
        self.assertLess(break_index, texts.index("상세보고서"))
        self.assertGreater(break_index, texts.index("1페이지 요약"))

    def test_bookmarks_and_internal_links(self):
        bookmarks = {element.get(qn("w:name")) for element in self.body.iter(qn("w:bookmarkStart"))}
        self.assertEqual(bookmarks, {"sec_overview", "sec_status", "sec_synergy", "sec_advice"})
        anchors = [element.get(qn("w:anchor")) for element in self.body.iter(qn("w:hyperlink"))]
        self.assertEqual(anchors, ["sec_overview", "sec_status", "sec_advice"])
        # 링크는 모두 실제 있는 책갈피를 가리켜야 해요.
        self.assertTrue(set(anchors) <= bookmarks)

    def test_footnote(self):
        footnote = next(text for text in docx_texts(self.content) if text.startswith("[1] "))
        self.assertIn("현대엘리베이터 · elv_2025_report.pdf · p.12", footnote)
        self.assertIn("“물류센터용 고속 화물 승강기", footnote)


class WatermarkTest(unittest.TestCase):

    def test_markdown_body_unchanged(self):
        body = report_exporter.to_markdown(make_report())
        marked = report_exporter.apply_watermark(body.encode("utf-8"), "md", WATERMARK).decode("utf-8")
        self.assertTrue(marked.startswith(f"> {WATERMARK}"))
        self.assertIn(body, marked)  # 본문·각주는 그대로 안에 들어 있어요
        self.assertTrue(marked.rstrip().endswith(WATERMARK))

    def test_docx_header_only(self):
        content = report_exporter.to_docx(make_report())
        marked = report_exporter.apply_watermark(content, "docx", WATERMARK)
        document = Document(io.BytesIO(marked))
        self.assertEqual(document.sections[0].header.paragraphs[0].text, WATERMARK)
        self.assertEqual(docx_texts(marked), docx_texts(content))  # 본문 문단은 그대로

    def test_unknown_format(self):
        with self.assertRaises(ValueError):
            report_exporter.apply_watermark(b"", "pdf", WATERMARK)


@patch(RECORD_AUDIT_LOG)
class ExportReportFileTest(unittest.TestCase):
    # 클래스 위 @patch는 모든 test_ 함수에 적용되고, 가짜 감사 로그가 audit 인자로 들어와요.

    def test_markdown_download_logged(self, audit):
        with patch(DECIDE, return_value=ALLOWED) as decide:
            content = report_exporter.export_report_file(USER, make_report(), "md")
        # A-06에는 각주의 문서 ID를 순서대로 중복 없이, 보고서 ID는 글자로, 최고 보안 등급도 같이 넘겨요.
        decide.assert_called_once_with(
            USER, ["doc-elv", "doc-mvx", "doc-hri"], "md", "12", max_access_level="public",
        )
        # 워터마크는 A-06이 준 문구를 그대로 넣어요.
        self.assertTrue(content.decode("utf-8").startswith(f"> {WATERMARK}"))
        audit.assert_called_once_with(
            7, "download", "12", None, False, None,
            target_type="report", doc_ids=["doc-elv", "doc-mvx", "doc-hri"],
            doc_companies=["ELV", "MVX", "HRI"], context_detail={"file_format": "md"},
        )

    def test_docx_download(self, audit):
        with patch(DECIDE, return_value=ALLOWED) as decide:
            content = report_exporter.export_report_file(USER, make_report(report_id=None), "docx")
        self.assertIn("스마트 물류 종합 보고서", docx_texts(content))
        self.assertIsNone(decide.call_args[0][3])  # 저장 전 보고서는 report_id 없음
        self.assertIsNone(audit.call_args[0][2])

    def test_blocked_raises_without_second_log(self, audit):
        with patch(DECIDE, return_value=BLOCKED), \
             patch.object(report_exporter, "to_docx") as to_docx, \
             self.assertRaises(PermissionError) as raised:
            report_exporter.export_report_file(USER, make_report(), "docx")
        self.assertIn("열람 권한이 없는 근거 문서", str(raised.exception))
        to_docx.assert_not_called()  # 차단이면 파일을 만들지 않아요
        audit.assert_not_called()    # 차단 기록은 A-06이 남겨서 R-05는 또 남기지 않아요

    def test_no_watermark(self, audit):
        with patch(DECIDE, return_value={**ALLOWED, "watermark": None}):
            content = report_exporter.export_report_file(USER, make_report(), "md")
        self.assertTrue(content.decode("utf-8").startswith("# 스마트 물류 종합 보고서"))

    def test_audit_failure_does_not_block(self, audit):
        audit.side_effect = RuntimeError("DB 연결 실패")
        with patch(DECIDE, return_value=ALLOWED), \
             self.assertLogs("app.reports.services.report_exporter", level="ERROR"):
            content = report_exporter.export_report_file(USER, make_report(), "md")
        self.assertIsNotNone(content)

    def test_unknown_format(self, audit):
        with patch(DECIDE) as decide, self.assertRaises(ValueError):
            report_exporter.export_report_file(USER, make_report(), "pdf")
        decide.assert_not_called()
        audit.assert_not_called()


@patch(DECIDE, return_value=ALLOWED)
@patch(RECORD_AUDIT_LOG)
class ExportReportTest(unittest.TestCase):
    """명세서 export_report(user, report_id, file_format): 저장된 본인 보고서를 찾아 내보낸다."""

    def test_own_report(self, _audit, _decide):
        with patch(FIND_BY_ID, return_value=make_report()) as find:
            content = report_exporter.export_report(USER, 12, "md")
        find.assert_called_once_with(12, 7)  # 보고서 ID + 로그인 사용자 ID
        self.assertIn("# 스마트 물류 종합 보고서", content.decode("utf-8"))

    def test_not_found(self, _audit, decide):
        # 남의 보고서거나 없는 보고서면 find_by_id가 None → 파일을 만들지 않아요.
        with patch(FIND_BY_ID, return_value=None), self.assertRaises(LookupError):
            report_exporter.export_report(USER, 99, "md")
        decide.assert_not_called()


def _buttons_app():
    """AppTest로 띄울 화면: 다운로드 버튼 + 화면을 다시 그리게 하는 아무 버튼."""
    import streamlit as st

    from app.reports.views.report_view import render_export_buttons
    from tests.reports.test_report_exporter import USER, make_report

    render_export_buttons(USER, make_report())
    st.button("다른 버튼", key="other")


@patch(DECIDE, return_value=BLOCKED)
@patch(RECORD_AUDIT_LOG)
class ExportButtonsTest(unittest.TestCase):
    """화면을 그리기만 해서는 A-06 판정(=차단 감사 로그)이 불리지 않는지 확인한다."""

    def test_render_does_not_decide(self, audit, decide):
        from streamlit.testing.v1 import AppTest

        app = AppTest.from_function(_buttons_app, default_timeout=15)
        app.run()
        app.button(key="other").click().run()  # 화면을 한 번 더 그려도
        self.assertFalse(app.exception)
        decide.assert_not_called()  # 누르기 전에는 판정하지 않아요 → 차단 기록이 쌓이지 않아요
        audit.assert_not_called()


class FileNameTest(unittest.TestCase):

    def test_file_name(self):
        self.assertEqual(report_exporter.build_file_name(make_report(), "md"), "스마트_물류_종합_보고서.md")
        report = make_report()
        report.request = type(MOCK_REQUEST)(**{**MOCK_REQUEST.__dict__, "topic": 'AI/로봇: "협업"?'})
        self.assertEqual(report_exporter.build_file_name(report, "docx"), "AI_로봇_협업_종합_보고서.docx")


if __name__ == "__main__":
    unittest.main()
